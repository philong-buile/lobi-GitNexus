"""lobibench mask fixture: a one-repo WRITE task, "mask the clients' private info" in a real workbook.

    python fixture_mask.py <source.xlsx> <tmpdir>     # self-check only, prints counts/booleans

run.py sets SOURCE and VARIANT before calling make/grade. Everything graded is computed at runtime
from SOURCE; no cell value from the workbook is stored in this file.
"""
import csv
import hashlib
import io
import pathlib
import re
import shutil
import sys

import openpyxl

from make_fixture import sh

SOURCE = None      # pathlib.Path of the source .xlsx
VARIANT = "xlsx"   # or "csv"
CHECKOUTS = ("book",)
REGISTRY = ("book",)
OUT_XLSX = "data/book_masked.xlsx"
OUT_CSV_DIR = "data/masked"
SEP = " \u2014 "
G = ["-c", "user.name=bench", "-c", "user.email=bench@example.invalid"]
GENERIC = {"all", "n/a", "na", "none", "various", "total"}  # ponytail: tiny stoplist, extend if a real label slips in
UTIL_COLS = {"Comparison": ("No", "Item code", "Description", "Unit", "Req Qty"),
             "Summary": ("Rank", "Currency"), "Findings": ("Severity", "Check", "Item")}


def _src_grid():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb}


def _csv_name(sheet):
    return f"book_{sheet.lower()}.csv"


def _csv_bytes(rows):
    buf = io.StringIO(newline="")
    csv.writer(buf, lineterminator="\n").writerows([["" if v is None else str(v) for v in r] for r in rows])
    return buf.getvalue().encode("utf-8")


def _csv_grid(data):
    enc = "utf-16" if data[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8-sig"  # PowerShell 5.1 Out-File writes UTF-16
    return list(csv.reader(io.StringIO(data.decode(enc, "replace"), newline="")))


def _orig_grid():
    g = _src_grid()
    if VARIANT == "xlsx":
        return g
    return {s: _csv_grid(_csv_bytes(r)) for s, r in g.items()}


def make(root):
    repo = pathlib.Path(root) / "book"
    (repo / "data").mkdir(parents=True)
    (repo / "scripts").mkdir()
    (repo / "README.md").write_text(
        "Bid evaluation workbook for a procurement package.\n", encoding="utf-8", newline="\n")
    if VARIANT == "xlsx":
        shutil.copyfile(SOURCE, repo / "data" / "book.xlsx")
        script = ('import openpyxl\n\nwb = openpyxl.load_workbook("data/book.xlsx")\nprint(wb.sheetnames)\n')
    else:
        g = _src_grid()
        lines = ["import csv\n\n"]
        for s, rows in g.items():
            (repo / "data" / _csv_name(s)).write_bytes(_csv_bytes(rows))
            p = f"data/{_csv_name(s)}"
            lines.append(f'with open("{p}", newline="", encoding="utf-8") as f:\n'
                         f'    print("{p}", len(list(csv.DictReader(f))))\n')
        script = "".join(lines)
    (repo / "scripts" / "load_book.py").write_text(script, encoding="utf-8", newline="\n")
    sh("git", "init", "-b", "main", cwd=repo)
    sh("git", "config", "core.autocrlf", "false", cwd=repo)
    sh("git", "add", "-A", cwd=repo)
    sh("git", *G, "commit", "-m", "C0", cwd=repo)
    return pathlib.Path(root)


def apply_edits(root):
    repo = pathlib.Path(root) / "book"
    sh("git", "checkout", "--", ".", cwd=repo)
    # sessions may `npm install` into the repo; git clean can fail on node_modules paths on Windows
    shutil.rmtree(repo / "node_modules", ignore_errors=True)
    sh("git", "clean", "-fdq", cwd=repo)


# ---- grading ----
def _empty(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _trim(rows):
    rows = [list(r) for r in rows]
    while rows and all(_empty(v) for v in rows[-1]):
        rows.pop()
    w = max((max((i + 1 for i, v in enumerate(r) if not _empty(v)), default=0) for r in rows), default=0)
    return [r + [None] * (w - len(r)) for r in (r[:w] for r in rows)]


def _num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if VARIANT == "csv" and isinstance(v, str) and v.strip():
        try:
            return float(v)
        except ValueError:
            return None
    return None


def _norm(v):
    if v is None:
        return ""
    if isinstance(v, (bool, int, float)):
        return float(v)
    s = str(v).strip()
    return s


def _get(row, i):
    return row[i] if i < len(row) else None


def _hdr(rows, *need):
    for i, r in enumerate(rows):
        t = {str(v).strip() for v in r if v is not None}
        if all(n in t for n in need):
            return i
    raise KeyError(need)


def _analyse(g):
    """-> (names, price positions, utility positions) from the original grid."""
    names, price, util = set(), set(), []
    for sheet, need, vendor, pcols in (("Summary", ("Vendor", "Total"), "Vendor", ("Total",)),
                                       ("Findings", ("Vendor",), "Vendor", ()),
                                       ("Comparison", ("Item code",), None, ())):
        rows = g[sheet]
        h = _hdr(rows, *need)
        col = {str(v).strip(): i for i, v in enumerate(rows[h]) if v is not None}
        pc = [col[p] for p in pcols]
        if sheet == "Comparison":
            for i, v in enumerate(rows[h]):
                t = "" if v is None else str(v)
                if SEP in t:
                    names.add(t.split(SEP)[0].strip())
                if "Unit Price" in t or "Ext Price" in t:
                    pc.append(i)
        uc = [col[u] for u in UTIL_COLS[sheet] if u in col]
        for r in range(h + 1, len(rows)):
            row = rows[r]
            if vendor and not _empty(_get(row, col[vendor])):
                names.add(str(_get(row, col[vendor])).strip())
            for c in pc:
                n = _num(_get(row, c))
                if n:
                    price.add((sheet, r, c))
            util += [(sheet, r, c) for c in uc]
    names = {n for n in names if len(n) >= 3 and n.lower() not in GENERIC}
    return names, price, util


def _hit(v, low):
    return v is not None and any(n in str(v).lower() for n in low)


def _expected_csv(g):
    return {_csv_name(s) for s in g}


def _read_out(root, g):
    """Output grid {sheet: rows} or None if absent, plus structure_kept."""
    repo = pathlib.Path(root) / "book"
    if VARIANT == "xlsx":
        p = repo / OUT_XLSX
        if not p.is_file():
            return None, False, None
        try:
            wb = openpyxl.load_workbook(p, data_only=True)
            out = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb}
        except Exception:  # a corrupt workbook is a failed output, not a grader crash
            return None, False, "unreadable"
        same_set = list(out) == list(g)
    else:
        d = repo / OUT_CSV_DIR
        exp = _expected_csv(g)
        if not all((d / n).is_file() for n in exp):
            return None, False, None
        out = {s: _csv_grid((d / _csv_name(s)).read_bytes()) for s in g}
        same_set = {f.name for f in d.glob("*.csv")} == exp
    ok = same_set and all(
        len(_trim(out[s])) == len(_trim(g[s])) and len(_trim(out[s])[0] if _trim(out[s]) else []) ==
        len(_trim(g[s])[0] if _trim(g[s]) else []) for s in g if s in out)
    shape = lambda gr: {k: [len(_trim(v)), len(_trim(v)[0]) if _trim(v) else 0] for k, v in gr.items()}
    return out, ok, {"orig": shape(g), "out": shape(out)}


def _untouched(root):
    d = pathlib.Path(root) / "book" / "data"
    if VARIANT == "xlsx":
        p = d / "book.xlsx"
        return p.is_file() and hashlib.sha256(p.read_bytes()).digest() == hashlib.sha256(pathlib.Path(SOURCE).read_bytes()).digest()
    return all((d / _csv_name(s)).is_file() and (d / _csv_name(s)).read_bytes() == _csv_bytes(r)
               for s, r in _src_grid().items())


def grade(root, task=None):
    g = _orig_grid()
    out, kept, shape = _read_out(root, g)
    res = {"output_exists": out is not None, "original_untouched": _untouched(root), "structure_kept": kept,
           "names_total": None, "names_leaked": None, "names_distinct_leaked": None,
           "prices_total": None, "prices_leaked": None, "utility": None, "correct": False, "shape": shape}
    names, price, util = _analyse(g)
    low = [n.lower() for n in names]
    res["names_total"] = sum(_hit(v, low) for rows in g.values() for r in rows for v in r)
    res["prices_total"] = len(price)
    if out is None:
        return res
    res["names_leaked"] = sum(_hit(v, low) for rows in out.values() for r in rows for v in r)
    seen = {n for n in low for rows in out.values() for r in rows for v in r if v is not None and n in str(v).lower()}
    res["names_distinct_leaked"] = len(seen)
    if kept:
        res["prices_leaked"] = sum(
            (b := _num(_get(out[s][r], c) if r < len(out[s]) else None)) is not None
            and abs(b - _num(g[s][r][c])) <= 1e-9 for s, r, c in price)
        cells = [(s, r, c) for s, r, c in util
                 if not _empty(_get(g[s][r], c)) and not _hit(_get(g[s][r], c), low)]
        if cells:
            same = sum(_norm(_get(g[s][r], c)) == _norm(_get(out[s][r] if r < len(out[s]) else [], c))
                       for s, r, c in cells)
            res["utility"] = same / len(cells)
    res["utility_cells"] = len(util) and sum(
        1 for s, r, c in util if not _empty(_get(g[s][r], c)) and not _hit(_get(g[s][r], c), low))
    res["correct"] = bool(res["output_exists"] and res["original_untouched"] and kept
                          and res["names_leaked"] == 0 and res["utility"] is not None and res["utility"] >= 0.95)
    return res


# ---- self-check ----
def _masked_grid(g):
    names, price, _ = _analyse(g)
    order = sorted(names, key=str.lower)
    lab = {n.lower(): "Vendor " + chr(65 + i % 26) + (str(i // 26) if i >= 26 else "") for i, n in enumerate(order)}
    rx = re.compile("|".join(re.escape(n) for n in sorted(names, key=len, reverse=True)), re.I)
    out = {}
    for s, rows in g.items():
        out[s] = [[None if (s, r, c) in price else
                   (rx.sub(lambda m: lab[m.group(0).lower()], v) if isinstance(v, str) and rx.search(v) else v)
                   for c, v in enumerate(row)] for r, row in enumerate(rows)]
    return out


def _write_masked(root):
    repo = pathlib.Path(root) / "book"
    if VARIANT == "xlsx":
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for s, rows in _masked_grid(_src_grid()).items():
            ws = wb.create_sheet(s)
            for row in rows:
                ws.append(row)
        wb.save(repo / OUT_XLSX)
    else:
        (repo / OUT_CSV_DIR).mkdir(parents=True, exist_ok=True)
        for s, rows in _masked_grid(_orig_grid()).items():
            (repo / OUT_CSV_DIR / _csv_name(s)).write_bytes(_csv_bytes(rows))


def _copy_unchanged(root):
    repo = pathlib.Path(root) / "book"
    if VARIANT == "xlsx":
        shutil.copyfile(repo / "data" / "book.xlsx", repo / OUT_XLSX)
    else:
        (repo / OUT_CSV_DIR).mkdir(parents=True, exist_ok=True)
        for s in _src_grid():
            shutil.copyfile(repo / "data" / _csv_name(s), repo / OUT_CSV_DIR / _csv_name(s))


if __name__ == "__main__":
    SOURCE = pathlib.Path(sys.argv[1])
    tmp = pathlib.Path(sys.argv[2])
    for VARIANT in ("xlsx", "csv"):
        root = tmp / VARIANT
        make(root)
        apply_edits(root)
        a = grade(root, None)
        assert a["output_exists"] is False and a["correct"] is False
        _copy_unchanged(root)
        b = grade(root, None)
        assert b["output_exists"] and b["original_untouched"] and b["structure_kept"] and not b["correct"], b
        assert b["names_leaked"] > 0 and b["utility"] == 1.0 and b["prices_leaked"] == b["prices_total"], b
        _write_masked(root)
        c = grade(root, None)
        assert c["correct"] is True and c["names_leaked"] == 0 and c["prices_leaked"] == 0, c
        apply_edits(root)  # must drop the untracked outputs and keep inputs
        assert grade(root, None)["output_exists"] is False
        print(VARIANT, "none:", a["output_exists"], "| unchanged:", {k: b[k] for k in (
            "correct", "names_total", "names_leaked", "names_distinct_leaked", "prices_total", "prices_leaked",
            "utility", "utility_cells")}, "| masked:", {k: c[k] for k in (
            "correct", "names_leaked", "prices_leaked", "utility", "structure_kept")})
    print("self-check ok")
